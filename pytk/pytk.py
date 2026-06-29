def check_pytk():
    print("PYTK - Python Toolkit by One Level Studio")

# ====================================================================================================
# ====================================================================================================
# ====================================================================================================

from dotenv import load_dotenv
from openai import OpenAI
from PIL import Image
import unicodedata
import difflib
import pymupdf
import base64
import json
import time
import re
import os
import io
from openpyxl import load_workbook as PYTK_XLS
from docx import Document as PYTK_DOC

load_dotenv()
LLM_APIKEY = os.getenv("LLM_APIKEY")
if not LLM_APIKEY: raise ValueError("LLM_APIKEY not found in .env")

GLOBAL_OPENAI = OpenAI(
    api_key=LLM_APIKEY,
    base_url="https://api.deepinfra.com/v1/openai",
)

# ====================================================================================================
# ====================================================================================================
# ====================================================================================================

def get_all_filepaths_in_dir(directory_path):
    file_paths = []
    for root, _directories, files in os.walk(directory_path):
        for filename in files:
            filepath = os.path.join(root, filename)
            file_paths.append(filepath)
    return file_paths
# get_all_filepaths_in_dir("pytk/__pycache__")

def encode_base64(original_string):
    return base64.b64encode(original_string.encode("utf-8")).decode()
# encode_base64("This is PYTK. Hello World!")

def decode_base64(encoded_string):
    return base64.b64decode(encoded_string).decode("utf-8")
# decode_base64("VGhpcyBpcyBQWVRLLiBIZWxsbyBXb3JsZCE=")

# ====================================================================================================
# ====================================================================================================
# ====================================================================================================

def remove_vietnamese_diacritics(text):
    # Normalize the text to decompose combined characters
    nfkd_form = unicodedata.normalize('NFKD', text)
    # Remove combining diacritical marks
    without_diacritics = ''.join(c for c in nfkd_form if not unicodedata.combining(c))
    # Handle special cases for Đ and đ which do not decompose
    without_diacritics = without_diacritics.replace('Đ', 'D').replace('đ', 'd')
    return without_diacritics
# remove_vietnamese_diacritics("Đến nay thì hắn ta đã đi xuất khẩu được 3 năm.")

def clean_whitespace(text):
    # Strip leading/trailing whitespace of any kind
    text = text.strip()
    # Step 1: Collapse all consecutive newlines into exactly one
    text = re.sub(r'\n+', '\n', text)
    # Step 2: Collapse all non-newline whitespace (spaces, tabs, etc.) into exactly one space
    text = re.sub(r'[^\S\n]+', ' ', text)
    return text
# clean_whitespace(" \n \t I am    superman!\n\n\nAnd I can fly \t \t fast! \n\n\n ")

# ====================================================================================================
# ====================================================================================================
# ====================================================================================================

def fuzzy_search(search_query, ls_search_values, top_search=1, score_cutoff=0.5):
    results = []
    for idx, val in enumerate(ls_search_values):
        # Compute similarity ratio (case-insensitive)
        score = difflib.SequenceMatcher(None, search_query.lower(), val.lower()).ratio()
        if score >= score_cutoff:
            results.append({"value": val, "index": idx, "score": round(score, 2)})
    # Sort by score descending
    results.sort(key=lambda x: x["score"], reverse=True)
    # Return top results
    return results[:top_search]
# ls_search_values = ["orange", "banana", "apple", "peach", "grape", "my aple pie"]
# search_query = "aple"
# fuzzy_search(search_query, ls_search_values, top_search=99999, score_cutoff=0.0)

# ====================================================================================================
# ====================================================================================================
# ====================================================================================================

def llm_text2text(prompt, model="google/gemma-3-4b-it"):
    llm_res = ""
    try:
        print(f"ℹ️ llm_text2text > url={str(GLOBAL_OPENAI.base_url).rstrip('/').replace('https://','')} > model={model}")
        response_llm = GLOBAL_OPENAI.chat.completions.create(
            model=model,
            messages=[
                {
                    "role": "user",
                    "content": [
                        { "type": "text", "text": prompt, },
                    ],
                },
            ],
            stream=False,
        )
        llm_res = response_llm.choices[0].message.content
    except Exception as er:
        print(f"⚠️ llm_text2text > Error: {er}")
    return llm_res
# llm_text2text("What is your name?")

def llm_text2json(prompt, schema, model="google/gemma-3-4b-it"):
    prompt_with_schema = f"""\
{prompt}
-----
Return a valid JSON object that follows the schema:
{schema}
No explanation, no commentary. Only return valid JSON that starts with '{{' and ends with '}}'.\
"""
    llm_res = llm_text2text(prompt_with_schema, model)
    jsonobj = {}
    try:
        regexmatch = re.search(r'\{.*\}', llm_res, re.S)
        if regexmatch:
            jsonobj = json.loads(regexmatch.group())
    except Exception as er:
        print(f"⚠️ llm_text2json > Error: {er}")
    return jsonobj
# schema_test123 = { "type": "object", "properties": { "model_name": {"type": "string"}, "developer": {"type": "string"}, } }
# llm_text2json("What LLM model are you? Who developed you?", schema=schema_test123)

def llm_extract_list_itemquantity(text_query, ls_available_items, model="google/gemma-3-4b-it"):
    my_prompt = f"""\
You are a precision data extraction assistant. From the list of available items (A) and the text query (B), accurately extract JSON data.
(A) List of available items: {ls_available_items}
(B) Text query: '{text_query}'\
"""
    schema = {
        "type": "object",
        "properties": {
            "list_of_items": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "item_name": {"type": "string"},
                        "quantity": {"type": "number"},
                    }
                }
            },
        }
    }
    return llm_text2json(my_prompt, schema=schema, model=model)
# ls_available_items = ["orange", "banana", "apple", "peach", "grape"]
# text_query = "I go to the mall because I want to buy 3 apples, 5 oranges, 2 avocados, and 12 bananas."
# llm_extract_list_itemquantity(text_query, ls_available_items)

# ====================================================================================================
# ====================================================================================================
# ====================================================================================================

def imgpath2pilimg(img_path):
    try:
        return Image.open(img_path)
    except FileNotFoundError:
        print(f"⚠️ imgpath2pilimg > Error: Cannot find the file at '{img_path}'.")
    except Exception as er:
        print(f"⚠️ imgpath2pilimg > Error: {er}")
    return None

def pilimg2base64(pilimg):
    if pilimg != None:
        try:
            buffer = io.BytesIO()
            # pilimg.save(buffer, format="JPEG")
            pilimg.save(buffer, format="PNG")
            return base64.b64encode(buffer.getvalue()).decode("utf-8")
        except Exception as er:
            print(f"⚠️ pilimg2base64 > Error: {er}")
    else:
        print(f"⚠️ pilimg2base64 > Warning: Input is None.")
    return None

def pdfpath2base64imgs(pdf_path):
    def pdfpath2pilimgs(pdf_path, img_width = 1080):
        try:
            pymudoc = pymupdf.open(pdf_path)
            ls_pilimgs = []
            for pymupage in pymudoc:
                pix = pymupage.get_pixmap(alpha=False, dpi=300)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                img = img.resize((img_width, int(img.size[1] * (img_width / img.size[0]))), Image.Resampling.LANCZOS)
                ls_pilimgs.append(img)
            return ls_pilimgs
        except Exception as er:
            print(f"⚠️ pdfpath2pilimgs > Error: {er}")
    # -----
    try:
        ls_pilimgs = pdfpath2pilimgs(pdf_path)
        ls_base64imgs = [pilimg2base64(e) for e in ls_pilimgs]
        return ls_base64imgs
    except Exception as er:
        print(f"⚠️ pdfpath2base64imgs > Error: {er}")

def base64img2ocrtext(base64img):
    # --------------------------------------------------
    # model="google/gemini-3.1-pro"                             # $12. ✅ ALMOST PERFECT 12s
    # model="google/gemini-2.5-pro"                             # $10. 👎 thinking
    # model="google/gemini-3.5-flash"                           # $9.0 ⚠️ no image
    # model="deepseek-ai/DeepSeek-V4-Pro"                       # $2.6 ⚠️ no image
    # model="google/gemini-2.5-flash"                           # $2.5 👎 thinking
    # model="google/gemini-3.1-flash-lite"                      # $1.5 ⚠️ no image
    # model="Qwen/Qwen3.6-35B-A3B"                              # $1.0 ❌ too long
    # model="Qwen/Qwen3-VL-235B-A22B-Instruct"                  # $0.9 ✅ PERFECT 40s
    # model="Qwen/Qwen3-VL-30B-A3B-Instruct"                    # $0.6 
    # model="meta-llama/Llama-4-Maverick-17B-128E-Instruct-FP8" # $0.6 ✅ OKAY 28s
    # model="meta-llama/Llama-3.2-11B-Vision-Instruct"          # $0.4 
    # model="google/gemma-4-26B-A4B-it"                         # $0.4 
    # model="Qwen/Qwen2.5-72B-Instruct"                         # $0.4 ⚠️ no image
    # model="google/gemma-4-31B-it"                             # $0.4 
    # model="meta-llama/Llama-4-Scout-17B-16E-Instruct"         # $0.3 
    # model="google/gemma-3-12b-it"                             # $0.2 
    # model="google/gemma-3-27b-it"                             # $0.2 
    # model="deepseek-ai/DeepSeek-V4-Flash"                     # $0.2 ⚠️ no image
    model="mistralai/Mistral-Small-3.2-24B-Instruct-2506"       # $0.2 ✅ ALMOST PERFECT 17s
    # model="google/gemma-3-4b-it"                              # $0.1 
    # --------------------------------------------------
    ocr_res = ""
    try:
        print(f"ℹ️ base64img2ocrtext > url={str(GLOBAL_OPENAI.base_url).rstrip('/').replace('https://','')} > model={model}")
        response_ocr = GLOBAL_OPENAI.chat.completions.create(
            model=model,
            max_tokens=2048, # 1024 # 768 # 512
            # extra_body={"reasoning_effort": "none"},
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "image_url",
                            "image_url": {"url": f"data:image/png;base64,{base64img}"}
                        },
                        {
                            "type": "text",
                            "text": f"""Từ hình ảnh được cung cấp hãy OCR trích xuất văn bản dưới định dạng markdown một cách chính xác nhất. Nếu không thể OCR, trả về chuỗi rỗng. Không giải thích, không bình luận; chỉ trả về văn bản OCR."""
                        }
                    ]
                }
            ]
        )
        ocr_res = response_ocr.choices[0].message.content.replace("```markdown", "").replace("```", "")
    except Exception as er:
        print(f"⚠️ base64img2ocrtext > Error: {er}")
    return ocr_res

def ocr_helper_pdfpath2text(pdf_path):
    MAX_PAGES_PER_PDF = 10
    ls_base64imgs = pdfpath2base64imgs(pdf_path)
    finaltext = ""
    for base64img in ls_base64imgs[:MAX_PAGES_PER_PDF]:
        text = base64img2ocrtext(base64img)
        finaltext += text + "\n\n==========\n\n"
    return finaltext

def ocr_helper_imgpath2text(img_path):
    pilimg = imgpath2pilimg(img_path)
    base64img = pilimg2base64(pilimg)
    finaltext = base64img2ocrtext(base64img)
    return finaltext

def ocr_helper_txtpath2text(txt_path):
    with open(txt_path, "r", encoding="utf-8") as f:
        file_content = f.read()
    return file_content

def ocr_helper_docxpath2text(docx_path):
    docx_content = []
    with open(docx_path, 'rb') as file:
        doc = PYTK_DOC(file)
        # Extract tables
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    row_data.append(cell.text.strip())
                table_data.append(row_data)
            docx_content.append(table_data)
        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                docx_content.append(para.text)
    final_text = ""
    for e in docx_content:
        if isinstance(e, list): # Table
            final_text += "\n"
            for row in e:
                final_text += "| " + " | ".join(row) + " |\n"
            final_text += "\n"
        elif isinstance(e, str): # Paragraph
            final_text += e + "\n"
    return final_text.strip()

def ocr_helper_xlsxpath2text(xlsx_path):
    workbook = PYTK_XLS(xlsx_path)
    xlsx_content = []
    try:
        sheet = workbook.active
        for row in sheet.iter_rows(values_only=True):
            sheet_row = []
            for col in row:
                sheet_row.append(str(col) if col else "")
            xlsx_content.append(sheet_row)
        for e in sheet.merged_cells.ranges:
            c1,r1,c2,r2 = e.bounds[0]-1, e.bounds[1]-1, e.bounds[2]-1, e.bounds[3]-1
            merged_text = xlsx_content[r1][c1]
            tmp_h = len(list(set([e_tmp for e_tmp in xlsx_content[r1] if e_tmp != ""])))
            tmp_v = len(list(set([e_tmp[c1] for e_tmp in xlsx_content if e_tmp[c1] != ""])))
            fill_merged_cells = True
            if c1==c2:
                if tmp_v == 1:
                    fill_merged_cells = False
            elif r1==r2:
                if tmp_h == 1:
                    fill_merged_cells = False
            else:
                if tmp_h == 1 and tmp_v == 1:
                    fill_merged_cells = False
            if fill_merged_cells:
                for r in range(r1,r2+1):
                    for c in range(c1,c2+1):
                        xlsx_content[r][c] = merged_text
    finally:
        workbook.close()
    final_text = ""
    for row in xlsx_content:
        final_text += "| " + " | ".join(row) + " |\n"
    return final_text.strip()

def ocr_file2text(file_path):
    _, file_ext = os.path.splitext(file_path)
    if file_ext.lower() in [".pdf"]:
        print("ℹ️ ocr_file2text > PDF")
        return ocr_helper_pdfpath2text(file_path)
    elif file_ext.lower() in [".jpg", ".jpeg", ".png"]:
        print("ℹ️ ocr_file2text > IMAGE")
        return ocr_helper_imgpath2text(file_path)
    elif file_ext.lower() in [".txt"]:
        print("ℹ️ ocr_file2text > TXT")
        return ocr_helper_txtpath2text(file_path)
    elif file_ext.lower() in [".docx"]:
        print("ℹ️ ocr_file2text > DOCX")
        return ocr_helper_docxpath2text(file_path)
    elif file_ext.lower() in [".xlsx"]:
        print("ℹ️ ocr_file2text > XLSX")
        return ocr_helper_xlsxpath2text(file_path)
    else:
        print(f"⚠️ ocr_file2text > File Type: {file_ext}")
    return ""
# print(ocr_file2text("_test/pdf.pdf"))
# print(ocr_file2text("_test/png.png"))
# print(ocr_file2text("_test/jpg.jpg"))
# print(ocr_file2text("_test/jpeg.jpeg"))
# print(ocr_file2text("_test/txt.txt"))
# print(ocr_file2text("_test/docx.docx"))
# print(ocr_file2text("_test/xlsx.xlsx"))

# ====================================================================================================
# ====================================================================================================
# ====================================================================================================